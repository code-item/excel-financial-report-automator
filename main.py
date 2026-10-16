import openpyxl

def manager_cost():
    product = input("Enter product or Exit: ")
    amount = int(input("Enter amount: "))
    type = input("Enter type or Exit: ")

    return {"product": product, "amount": amount, "type": type,}

list1 = []

while True:
    dict_0 = manager_cost()
    if dict_0.get("product") == "Exit":
        break
    if dict_0.get("type") == "Exit":
        break
    list1.append(dict_0)

amount = 0

for dict in list1:
    amount += dict.get("amount")

print(f"amount {amount}")

work_book = openpyxl.Workbook()

sheet = work_book.active

sheet.append(["Тип товара", "Название товара", "Количество конкретного товара"])

for dict in list1:
    sheet.append([dict.get("type"), dict.get("product"), dict.get("amount")])

row_text = sheet.max_row + 1
row_amount = sheet.max_row + 2

sheet[f"A{row_text}"] = "Суммарное количество товаров"
sheet.merge_cells(f"A{row_text}:C{row_text}")

sheet[f"A{row_amount}"] = amount
sheet.merge_cells(f"A{row_amount}:C{row_amount}")

work_book.save("Таблица.xlsx")